"""
Report Generator Service
========================
Generates formatted PDF compliance reports (via ReportLab) and editable Word reports (via python-docx)
summarising product compliance scan results, statutory declarations verification tables, and legal violations.

Supports:
  1. Single Scan PDF & DOCX Reports (accepts string UIDs like 'LM-2026-000030' or numeric IDs)
  2. Multi-Specimen Batch Compliance Audit Docket (consolidated executive summary + specimen audit sheets)
  3. On-the-fly PDF generation directly from in-memory scan objects

Public API
----------
    generate_pdf_report(scan_id: str | int, db=None) -> bytes
    generate_editable_report(scan_id: str | int, db=None) -> bytes
    generate_pdf_from_data(scan_data: dict[str, Any]) -> bytes
    generate_batch_pdf_report(scans_data: list[dict[str, Any]], batch_title: str, batch_id: str) -> bytes
"""

from __future__ import annotations

import io
import logging
from datetime import datetime, timezone
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

logger = logging.getLogger(__name__)

# Standard 7 statutory declarations under Legal Metrology Rules 2011
STATUTORY_DECLARATIONS_META = [
    {"key": "net_quantity", "label": "1. Net Quantity", "rule": "Rule 12 & Rule 6(1)(d)", "min_h": "4.00mm"},
    {"key": "mrp", "label": "2. Maximum Retail Price (MRP)", "rule": "Rule 6(1)(e)", "min_h": "2.50mm"},
    {"key": "manufacture_date", "label": "3. Date of Manufacture / PKD", "rule": "Rule 6(1)(f)", "min_h": "1.80mm"},
    {"key": "expiry_date", "label": "4. Date of Expiry / Best Before", "rule": "Rule 6(1)(f) & FSSAI", "min_h": "1.80mm"},
    {"key": "manufacturer_name_address", "label": "5. Manufacturer / Packer Address", "rule": "Rule 6(1)(a)", "min_h": "1.50mm"},
    {"key": "consumer_care_details", "label": "6. Consumer Care Contact Details", "rule": "Rule 6(1)(k)", "min_h": "1.50mm"},
    {"key": "country_of_origin", "label": "7. Country of Origin (Imported / Dual)", "rule": "Rule 6(10) & Rule 6(1)(b)", "min_h": "1.50mm"},
]


# ─────────────────────────────────────────────────────────────────────────────
# Normalization Helper
# ─────────────────────────────────────────────────────────────────────────────

def _normalize_scan_data(s: dict[str, Any]) -> dict[str, Any]:
    """Convert raw scan dictionary from ScanRepository or memory into standardized report structure."""
    scan_id = str(s.get("scan_id") or s.get("scan_uid") or "LM-UNKNOWN")
    raw_status = str(s.get("compliance_status") or s.get("status") or "non_compliant").lower().replace("-", "_")

    violations_raw = s.get("violations") or []
    violations = []
    for v in violations_raw:
        if isinstance(v, dict):
            violations.append({
                "field_name": v.get("field_name") or "Declaration",
                "violation_type": v.get("violation_type") or "unspecified",
                "severity": v.get("severity") or "medium",
                "details": v.get("details") or "Non-compliant declaration under PCR 2011",
                "rule_reference": v.get("rule_reference") or "Legal Metrology Rules 2011",
            })
        else:
            violations.append({
                "field_name": getattr(v, "field_name", "Declaration"),
                "violation_type": getattr(v.violation_type, "value", str(getattr(v, "violation_type", "unspecified"))),
                "severity": getattr(v.severity, "value", str(getattr(v, "severity", "medium"))),
                "details": getattr(v, "details", "N/A"),
                "rule_reference": getattr(v, "rule_reference", "Legal Metrology Rules 2011"),
            })

    if raw_status == "compliant" or (not violations and raw_status != "non_compliant"):
        verdict = "VERDICT: COMPLIANT"
        verdict_color = "green"
    elif any(v["violation_type"] in ["incorrect_format", "undersized_font", "missing"] for v in violations):
        verdict = "VERDICT: NON-COMPLIANT"
        verdict_color = "red"
    else:
        verdict = "VERDICT: PARTIAL REVIEW NEEDED"
        verdict_color = "amber"

    # Scale calibration
    scale_calib = s.get("scale_calibration") or {}
    calib_tier = scale_calib.get("tier") or scale_calib.get("calibration_tier") or "dpi_estimated"
    calib_label = scale_calib.get("label") or scale_calib.get("calibration_label") or "DPI Estimated (±0.5mm)"

    # Declarations fields dictionary
    fields = s.get("fields") or {}

    # Created at
    created_at = s.get("created_at") or datetime.now(timezone.utc)
    if isinstance(created_at, str):
        try:
            created_at = datetime.fromisoformat(created_at.replace("Z", "+00:00"))
        except Exception:
            created_at = datetime.now(timezone.utc)

    # Product name fallback
    prod_name = (
        s.get("product_name")
        or fields.get("brand_name", {}).get("extracted_value")
        or (s.get("product_details", {}).get("name") if isinstance(s.get("product_details"), dict) else None)
        or "Packaged Commodity Specimen"
    )

    return {
        "scan_id": scan_id,
        "scan_type": s.get("scan_type", "manual"),
        "scan_status": raw_status,
        "created_at": created_at,
        "product_id": s.get("product_id") or scan_id,
        "product_name": prod_name,
        "product_category": s.get("product_category") or s.get("category") or "Packaged Foods",
        "scanned_image_url": s.get("scanned_image_url") or "",
        "verdict": verdict,
        "verdict_color": verdict_color,
        "calibration_tier": calib_tier,
        "calibration_label": calib_label,
        "violations_count": len(violations),
        "violations": violations,
        "fields": fields,
        "expiry_intelligence": s.get("expiry_intelligence") or {},
        "damage_analysis": s.get("damage_analysis") or {},
    }


# ─────────────────────────────────────────────────────────────────────────────
# DB / Repository Fetch Helper
# ─────────────────────────────────────────────────────────────────────────────

async def _fetch_scan_data(scan_id: str | int, db: AsyncSession | None = None) -> dict[str, Any]:
    """
    Fetch Scan data from unified ScanRepository (SQLite), in-memory RECENT_SCANS,
    or PostgreSQL database.
    """
    str_id = str(scan_id).strip()

    # 1. Primary: Unified SQLite ScanRepository
    try:
        from app.core.scan_repository import ScanRepository
        repo_data = ScanRepository.get_scan(str_id)
        if repo_data:
            return _normalize_scan_data(repo_data)
    except Exception as exc:
        logger.debug("ScanRepository lookup error for %s: %s", str_id, exc)

    # 2. In-memory recent scans pool
    try:
        from app.routers.scans import RECENT_SCANS
        for s in RECENT_SCANS:
            if str(s.get("scan_id")) == str_id or str(s.get("scan_uid")) == str_id:
                return _normalize_scan_data(s)
    except Exception as exc:
        logger.debug("RECENT_SCANS lookup error: %s", exc)

    # 3. PostgreSQL Database lookup (if session is provided and ID is integer)
    if db is not None and str_id.isdigit():
        try:
            from app.models.scan import Scan
            int_id = int(str_id)
            stmt = select(Scan).where(Scan.id == int_id)
            res = await db.execute(stmt)
            scan_obj = res.scalar_one_or_none()
            if scan_obj:
                prod = scan_obj.product
                viols = list(scan_obj.violations or [])
                raw_dict = {
                    "scan_id": scan_obj.id,
                    "scan_type": getattr(scan_obj.scan_type, "value", str(scan_obj.scan_type)),
                    "compliance_status": getattr(scan_obj.status, "value", str(scan_obj.status)),
                    "created_at": scan_obj.created_at,
                    "product_name": prod.name if prod else "Packaged Commodity",
                    "product_category": prod.category if prod else "General",
                    "scanned_image_url": prod.scanned_image_url if prod else scan_obj.raw_image_url,
                    "violations": [
                        {
                            "field_name": v.field_name,
                            "violation_type": getattr(v.violation_type, "value", str(v.violation_type)),
                            "severity": getattr(v.severity, "value", str(v.severity)),
                            "details": v.details,
                            "rule_reference": "Legal Metrology Rules 2011",
                        }
                        for v in viols
                    ],
                }
                return _normalize_scan_data(raw_dict)
        except Exception as db_exc:
            logger.debug("PostgreSQL Scan table lookup error: %s", db_exc)

    # 4. Fallback: Demo / Test mock record
    if str_id in ["101", "102", "103"]:
        mock_data = {
            "scan_id": str_id,
            "product_name": "Premium Organic Almond Milk 1L",
            "product_category": "Beverages",
            "scan_type": "manual",
            "compliance_status": "non_compliant",
            "violations": [
                {
                    "field_name": "mrp",
                    "violation_type": "missing",
                    "severity": "high",
                    "details": "Mandatory declaration 'Maximum Retail Price (MRP)' is missing from label.",
                    "rule_reference": "Legal Metrology Rules 2011, Rule 6(1)(e)",
                },
                {
                    "field_name": "net_quantity",
                    "violation_type": "undersized_font",
                    "severity": "medium",
                    "details": "Font height 1.80mm is below required minimum 4.00mm for 1000g package size.",
                    "rule_reference": "Legal Metrology Rules 2011, Rule 7(1) Table I",
                },
            ],
            "fields": {
                "net_quantity": {"extracted_value": "1000 ml", "font_height_mm": 1.8},
                "mrp": {"extracted_value": None},
                "manufacture_date": {"extracted_value": "12/08/2026"},
                "expiry_date": {"extracted_value": "12/02/2027"},
                "manufacturer_name_address": {"extracted_value": "Organic Beverages Ltd, Plot 42, Mumbai 400001"},
                "consumer_care_details": {"extracted_value": "care@organicbev.in / 1800-200-1234"},
                "country_of_origin": {"extracted_value": "India"},
            },
        }
        return _normalize_scan_data(mock_data)

    raise ValueError(f"Scan record '{str_id}' was not found in the national compliance registry.")


# ─────────────────────────────────────────────────────────────────────────────
# 1. Single Scan PDF Report Builder (ReportLab)
# ─────────────────────────────────────────────────────────────────────────────

def _build_pdf_bytes(data: dict[str, Any]) -> bytes:
    """Build a comprehensive, publication-grade single scan PDF report."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()

    # Custom typography styles
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#1E293B"),
    )
    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#64748B"),
    )
    header_style = ParagraphStyle(
        "HeaderStyle",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#0F172A"),
        spaceBefore=8,
        spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "BodyStyle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#334155"),
    )
    bold_body = ParagraphStyle(
        "BoldBody",
        parent=body_style,
        fontName="Helvetica-Bold",
    )

    story = []

    # Title Banner & Header
    story.append(Paragraph("Legal Metrology Compliance Report", title_style))
    story.append(Paragraph(f"InnoveXguard AI Official Certification • Docket #{data['scan_id']}", subtitle_style))
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#CBD5E1"), spaceAfter=10))

    # Overall Verdict Banner
    verdict_colors_map = {
        "green": (colors.HexColor("#DCFCE7"), colors.HexColor("#15803D")),
        "red": (colors.HexColor("#FEE2E2"), colors.HexColor("#B91C1C")),
        "amber": (colors.HexColor("#FEF3C7"), colors.HexColor("#B45309")),
    }
    bg_c, text_c = verdict_colors_map.get(data["verdict_color"], (colors.HexColor("#F1F5F9"), colors.HexColor("#334155")))

    verdict_text_style = ParagraphStyle(
        "VerdictText",
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        textColor=text_c,
        alignment=1,
    )

    verdict_table = Table(
        [[Paragraph(f"{data['verdict']}", verdict_text_style)]],
        colWidths=[540],
    )
    verdict_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), bg_c),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BOX", (0, 0), (-1, -1), 1.2, text_c),
        ])
    )
    story.append(verdict_table)
    story.append(Spacer(1, 10))

    # Metadata Grid
    created_str = (
        data["created_at"].strftime("%Y-%m-%d %H:%M UTC")
        if isinstance(data["created_at"], datetime)
        else str(data["created_at"])
    )
    meta_table_data = [
        [
            Paragraph("<b>Product Name:</b>", body_style), Paragraph(str(data["product_name"]), body_style),
            Paragraph("<b>Docket / Scan ID:</b>", body_style), Paragraph(str(data["scan_id"]), bold_body),
        ],
        [
            Paragraph("<b>Category:</b>", body_style), Paragraph(str(data["product_category"]), body_style),
            Paragraph("<b>Scan Method:</b>", body_style), Paragraph(str(data["scan_type"]).upper(), body_style),
        ],
        [
            Paragraph("<b>Processed Date:</b>", body_style), Paragraph(created_str, body_style),
            Paragraph("<b>Compliance Status:</b>", body_style), Paragraph(str(data["scan_status"]).upper(), bold_body),
        ],
        [
            Paragraph("<b>Scale Calibration:</b>", body_style), Paragraph(str(data.get("calibration_label", "DPI Estimated (±0.5mm)")), bold_body),
            Paragraph("<b>Statutory Standard:</b>", body_style), Paragraph("Legal Metrology (PC) Rules, 2011", body_style),
        ],
    ]
    meta_table = Table(meta_table_data, colWidths=[95, 175, 95, 175])
    meta_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ])
    )
    story.append(Paragraph("1. Packaging & Specimen Metadata", header_style))
    story.append(meta_table)
    story.append(Spacer(1, 10))

    # 2. Mandatory Statutory Declarations Verification Table
    fields = data.get("fields") or {}
    decl_rows = [
        [
            Paragraph("<b>Statutory Declaration</b>", bold_body),
            Paragraph("<b>PCR 2011 Rule</b>", bold_body),
            Paragraph("<b>Extracted Value on Label</b>", bold_body),
            Paragraph("<b>Determination</b>", bold_body),
        ]
    ]

    for item in STATUTORY_DECLARATIONS_META:
        k = item["key"]
        val_obj = fields.get(k) or {}
        extracted_val = val_obj.get("extracted_value") if isinstance(val_obj, dict) else str(val_obj) if val_obj else None
        has_viol = any(v.get("field_name") == k for v in data.get("violations", []))

        if extracted_val:
            status_text = "<font color='#15803D'><b>VERIFIED PRESENT</b></font>" if not has_viol else "<font color='#B91C1C'><b>FLAGGED NON-COMPLIANT</b></font>"
            display_val = str(extracted_val)
        else:
            status_text = "<font color='#B91C1C'><b>MISSING</b></font>"
            display_val = "<i>Not declared on package</i>"

        decl_rows.append([
            Paragraph(item["label"], body_style),
            Paragraph(item["rule"], body_style),
            Paragraph(display_val, body_style),
            Paragraph(status_text, body_style),
        ])

    decl_table = Table(decl_rows, colWidths=[150, 110, 170, 110])
    decl_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ])
    )
    story.append(Paragraph("2. Mandatory Statutory Declarations Audit (Rule 6)", header_style))
    story.append(decl_table)
    story.append(Spacer(1, 10))

    # 3. Compliance Violations Detail Table
    story.append(Paragraph("3. Detected Statutory Infractions & Legal Citations", header_style))
    if not data["violations"]:
        story.append(
            Paragraph("🎉 <i>Zero statutory infractions detected. Packaging specimen satisfies mandatory Legal Metrology declarations.</i>", body_style)
        )
    else:
        v_rows = [
            [
                Paragraph("<b>#</b>", bold_body),
                Paragraph("<b>Field Name</b>", bold_body),
                Paragraph("<b>Violation Type</b>", bold_body),
                Paragraph("<b>Severity</b>", bold_body),
                Paragraph("<b>Evidence & Statutory Reference</b>", bold_body),
            ]
        ]
        for idx, v in enumerate(data["violations"], 1):
            sev = v["severity"].upper()
            sev_color = "#B91C1C" if sev in ["CRITICAL", "HIGH"] else "#B45309" if sev == "MEDIUM" else "#15803D"
            v_rows.append([
                Paragraph(str(idx), body_style),
                Paragraph(v["field_name"], bold_body),
                Paragraph(v["violation_type"].replace("_", " ").title(), body_style),
                Paragraph(f"<font color='{sev_color}'><b>{sev}</b></font>", body_style),
                Paragraph(f"{v['details']}<br/><font color='#64748B'><i>{v.get('rule_reference', '')}</i></font>", body_style),
            ])

        v_table = Table(v_rows, colWidths=[20, 100, 100, 60, 260])
        v_table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ])
        )
        story.append(v_table)

    story.append(Spacer(1, 14))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#E2E8F0"), spaceAfter=6))
    story.append(
        Paragraph(
            "<i>Evaluated under the Legal Metrology Act, 2009 & Packaged Commodities Rules, 2011 (as amended). "
            "Generated automatically by InnoveXguard AI. Certified for legal metrology administrative enforcement.</i>",
            subtitle_style,
        )
    )

    doc.build(story)
    return buffer.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# 2. Multi-Specimen Batch PDF Report Builder (ReportLab)
# ─────────────────────────────────────────────────────────────────────────────

def generate_batch_pdf_report(
    scans_data: list[dict[str, Any]],
    batch_title: str = "Multi-Specimen Batch Compliance Audit Docket",
    batch_id: str | None = None,
) -> bytes:
    """
    Generate a consolidated multi-page PDF compliance audit docket covering all specimens in a batch.
    Includes executive summary metrics, consolidated portfolio table, and individual product audit sheets.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()

    # Custom typography
    title_style = ParagraphStyle(
        "BatchTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0F172A"),
    )
    subtitle_style = ParagraphStyle(
        "BatchSub",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#64748B"),
    )
    h2_style = ParagraphStyle(
        "H2Style",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11.5,
        leading=15,
        textColor=colors.HexColor("#0F172A"),
        spaceBefore=10,
        spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "BatchBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#334155"),
    )
    bold_body = ParagraphStyle("BoldBatchBody", parent=body_style, fontName="Helvetica-Bold")

    story = []

    normalized_scans = [_normalize_scan_data(s) for s in scans_data]
    total_count = len(normalized_scans)
    compliant_count = sum(1 for s in normalized_scans if s["verdict_color"] == "green")
    non_compliant_count = sum(1 for s in normalized_scans if s["verdict_color"] == "red")
    partial_count = total_count - (compliant_count + non_compliant_count)
    total_violations = sum(s["violations_count"] for s in normalized_scans)
    compliance_rate = round((compliant_count / total_count * 100)) if total_count > 0 else 100

    docket_id = batch_id or f"DOCKET-2026-{datetime.now(timezone.utc).strftime('%m%d%H%M')}"
    timestamp_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    # ── Executive Cover Page ──────────────────────────────────────────────────
    story.append(Paragraph("Legal Metrology Multi-Specimen Batch Audit Docket", title_style))
    story.append(Paragraph(f"InnoveXguard AI • Consolidated Inspection Dossier • {batch_title}", subtitle_style))
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#1E293B"), spaceAfter=10))

    # Executive KPI Summary Grid
    kpi_data = [
        [
            Paragraph("<b>Batch Docket Ref:</b>", body_style), Paragraph(str(docket_id), bold_body),
            Paragraph("<b>Inspection Date:</b>", body_style), Paragraph(timestamp_str, body_style),
        ],
        [
            Paragraph("<b>Total Specimens:</b>", body_style), Paragraph(f"<b>{total_count}</b> packaging units", bold_body),
            Paragraph("<b>Overall Compliance:</b>", body_style),
            Paragraph(f"<b>{compliance_rate}% Compliance Rate</b>", bold_body),
        ],
        [
            Paragraph("<b>Compliant Units:</b>", body_style),
            Paragraph(f"<font color='#15803D'><b>{compliant_count} Verified</b></font>", body_style),
            Paragraph("<b>Non-Compliant Units:</b>", body_style),
            Paragraph(f"<font color='#B91C1C'><b>{non_compliant_count} Flagged</b></font>", body_style),
        ],
        [
            Paragraph("<b>Total Violations:</b>", body_style),
            Paragraph(f"<font color='#B91C1C'><b>{total_violations} Infractions</b></font>", body_style),
            Paragraph("<b>Regulatory Scope:</b>", body_style),
            Paragraph("PCR 2011 Rules 6, 7, 8, 12", body_style),
        ],
    ]
    kpi_table = Table(kpi_data, colWidths=[105, 165, 105, 165])
    kpi_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ])
    )
    story.append(Paragraph("1. Executive Batch Docket Summary", h2_style))
    story.append(kpi_table)
    story.append(Spacer(1, 12))

    # ── Consolidated Portfolio Ledger Table ───────────────────────────────────
    story.append(Paragraph("2. Specimen Compliance Ledger Table", h2_style))

    ledger_rows = [
        [
            Paragraph("<b>#</b>", bold_body),
            Paragraph("<b>Specimen ID</b>", bold_body),
            Paragraph("<b>Commodity Name</b>", bold_body),
            Paragraph("<b>Category</b>", bold_body),
            Paragraph("<b>Determination</b>", bold_body),
            Paragraph("<b>Infractions</b>", bold_body),
        ]
    ]

    for idx, s in enumerate(normalized_scans, 1):
        v_badge = (
            "<font color='#15803D'><b>COMPLIANT</b></font>"
            if s["verdict_color"] == "green"
            else "<font color='#B91C1C'><b>NON-COMPLIANT</b></font>"
            if s["verdict_color"] == "red"
            else "<font color='#B45309'><b>PARTIAL REVIEW</b></font>"
        )
        ledger_rows.append([
            Paragraph(str(idx), body_style),
            Paragraph(f"<b>{s['scan_id']}</b>", body_style),
            Paragraph(str(s["product_name"])[:40], body_style),
            Paragraph(str(s["product_category"]), body_style),
            Paragraph(v_badge, body_style),
            Paragraph(str(s["violations_count"]), bold_body),
        ])

    ledger_table = Table(ledger_rows, colWidths=[20, 85, 175, 110, 95, 55])
    ledger_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ])
    )
    story.append(ledger_table)
    story.append(Spacer(1, 14))

    # ── Individual Specimen Breakdown Audit Sheets ─────────────────────────────
    story.append(PageBreak())
    story.append(Paragraph("3. Detailed Specimen Audit Sheets", h2_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#CBD5E1"), spaceAfter=8))

    for idx, s in enumerate(normalized_scans, 1):
        specimen_flowables = []

        verdict_c = colors.HexColor("#15803D") if s["verdict_color"] == "green" else colors.HexColor("#B91C1C")
        bg_col = colors.HexColor("#DCFCE7") if s["verdict_color"] == "green" else colors.HexColor("#FEE2E2")

        header_text = f"Specimen #{idx}: {s['product_name']} ({s['scan_id']})"
        specimen_flowables.append(Paragraph(f"<b>{header_text}</b>", h2_style))

        # Mini banner
        banner_table = Table([[Paragraph(f"<b>{s['verdict']}</b> — {s['violations_count']} infractions flagged", bold_body)]], colWidths=[540])
        banner_table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), bg_col),
                ("BOX", (0, 0), (-1, -1), 0.8, verdict_c),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ])
        )
        specimen_flowables.append(banner_table)
        specimen_flowables.append(Spacer(1, 4))

        # Violations snippet
        if s["violations"]:
            v_items = []
            for v in s["violations"][:4]:
                sev_c = "#B91C1C" if v["severity"].upper() in ["CRITICAL", "HIGH"] else "#B45309"
                v_items.append([
                    Paragraph(f"• <b>{v['field_name']}</b>: {v['details']} (<font color='{sev_c}'>{v['severity'].upper()}</font>)", body_style)
                ])
            v_subtable = Table(v_items, colWidths=[540])
            v_subtable.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]))
            specimen_flowables.append(v_subtable)
        else:
            specimen_flowables.append(Paragraph("<i>All 7 statutory packaging declarations verified compliant.</i>", body_style))

        specimen_flowables.append(Spacer(1, 8))
        story.append(KeepTogether(specimen_flowables))

    # Official Certification Seal
    story.append(Spacer(1, 14))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#CBD5E1"), spaceAfter=8))
    cert_text = (
        "<b>CERTIFICATE OF BATCH AUDIT:</b> This multi-specimen compliance docket was compiled via the InnoveXguard AI "
        "inspection engine. Evaluated in accordance with Schedule II font tolerances, Rule 6 statutory declarations, "
        "and Rule 12 standard quantity brackets under the Legal Metrology Act, 2009."
    )
    story.append(Paragraph(cert_text, subtitle_style))

    doc.build(story)
    return buffer.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# 3. Editable Word Report Generator (python-docx)
# ─────────────────────────────────────────────────────────────────────────────

def _set_cell_background(cell, hex_color: str):
    """Set table cell background color in python-docx."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def _build_docx_bytes(data: dict[str, Any]) -> bytes:
    """Build editable Word (.docx) report using python-docx."""
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    title_p = doc.add_paragraph()
    title_run = title_p.add_run("Legal Metrology Compliance Report")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 41, 59)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run(f"InnoveXguard AI • Scan #{data['scan_id']}")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(10)
    sub_run.font.color.rgb = RGBColor(100, 116, 139)

    verdict_p = doc.add_paragraph()
    verdict_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    verdict_run = verdict_p.add_run(f"OVERALL VERDICT: {data['verdict']}")
    verdict_run.font.name = "Calibri"
    verdict_run.font.size = Pt(13)
    verdict_run.font.bold = True

    if data["verdict_color"] == "green":
        verdict_run.font.color.rgb = RGBColor(21, 128, 61)
    elif data["verdict_color"] == "red":
        verdict_run.font.color.rgb = RGBColor(185, 28, 28)
    else:
        verdict_run.font.color.rgb = RGBColor(180, 83, 9)

    doc.add_heading("1. Scan & Product Metadata", level=1)
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.style = "Table Grid"

    created_str = (
        data["created_at"].strftime("%Y-%m-%d %H:%M UTC")
        if isinstance(data["created_at"], datetime)
        else str(data["created_at"])
    )
    meta_pairs = [
        ("Product Name:", str(data["product_name"])),
        ("Scan ID:", str(data["scan_id"])),
        ("Category:", str(data["product_category"])),
        ("Scan Type:", str(data["scan_type"]).upper()),
        ("Processed Date:", created_str),
        ("Status:", str(data["scan_status"]).upper()),
        ("Scale Calibration:", str(data.get("calibration_label", "DPI Estimated (±0.5mm)"))),
        ("Legal Standard:", "Legal Metrology 2011 (Rule 6/7)"),
    ]

    for idx, (label, val) in enumerate(meta_pairs):
        r_idx = idx // 2
        c_idx = idx % 2
        cell = meta_table.cell(r_idx, c_idx)
        p = cell.paragraphs[0]
        r1 = p.add_run(f"{label} ")
        r1.font.bold = True
        p.add_run(val)

    doc.add_paragraph()
    doc.add_heading("2. Compliance Violations Detail", level=1)

    if not data["violations"]:
        p = doc.add_paragraph()
        r = p.add_run("No compliance violations detected. Label satisfies mandatory declarations!")
        r.font.italic = True
    else:
        v_table = doc.add_table(rows=1, cols=5)
        v_table.style = "Table Grid"
        hdr_cells = v_table.rows[0].cells
        headers = ["#", "Field Name", "Violation Type", "Severity", "Details"]

        for idx, text in enumerate(headers):
            cell = hdr_cells[idx]
            _set_cell_background(cell, "1E293B")
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

        for v_idx, v in enumerate(data["violations"], 1):
            row_cells = v_table.add_row().cells
            row_cells[0].paragraphs[0].text = str(v_idx)
            row_cells[1].paragraphs[0].text = v["field_name"]
            row_cells[2].paragraphs[0].text = v["violation_type"].replace("_", " ").title()
            row_cells[3].paragraphs[0].text = v["severity"].upper()
            row_cells[4].paragraphs[0].text = v["details"]

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# 4. Public API Functions
# ─────────────────────────────────────────────────────────────────────────────

async def generate_pdf_report(scan_id: str | int, db: AsyncSession | None = None) -> bytes:
    """Generate PDF report for scan_id as raw bytes."""
    data = await _fetch_scan_data(scan_id, db)
    return _build_pdf_bytes(data)


def generate_pdf_from_data(scan_data: dict[str, Any]) -> bytes:
    """Generate PDF report directly from a scan dictionary payload."""
    data = _normalize_scan_data(scan_data)
    return _build_pdf_bytes(data)


async def generate_editable_report(scan_id: str | int, db: AsyncSession | None = None) -> bytes:
    """Generate Word (.docx) editable report for scan_id as raw bytes."""
    data = await _fetch_scan_data(scan_id, db)
    return _build_docx_bytes(data)
